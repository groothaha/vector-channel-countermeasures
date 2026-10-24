"""report.md -> 공식 Apart 템플릿(.docx)에 채운 제출본. 템플릿 스타일(글꼴·제목·각주)을 그대로 쓴다.

제목·저자·초록은 템플릿 표의 칸에 넣고, 안내 상자와 안내 문단은 지운 뒤 본문을 템플릿 스타일로 붙인다.

    ~/docx_env/bin/python -I sprint/paper/build_docx.py TEMPLATE.docx OUT.docx
"""
import re
import sys
from pathlib import Path

import copy

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
TITLE = "Cutting the Wire: Countermeasures Against Learned Vector Channels Between Language Models"
AUTHOR, AFFIL = "Jiho\u00a0Yun", "Independent"  # 저자 칸이 좁아 이름이 두 줄로 꺾이지 않게 줄바꿈 없는 공백
FIG_WIDTH = {"fig1": 16.5, "fig2": 16.5, "fig3": 15.0, "fig4": 11.5, "fig5": 16.5}  # cm, 본문 폭 16.5cm
RPR = {}  # 템플릿 안내 문단에서 떼어 온 글자 서식(글꼴 Old Standard TT 등): "body", "Heading 2", "Heading 3"


def styled(run, kind="body"):
    """템플릿의 같은 종류 문단 런 서식을 복사해 입힌다. 기울임(안내 문구 표시)은 뺀다."""
    if kind in RPR:
        old = run._r.find(qn("w:rPr"))
        new = copy.deepcopy(RPR[kind])
        for tag in ("w:i", "w:iCs", "w:color", "w:highlight"):
            for e in new.findall(qn(tag)):
                new.remove(e)
        if old is not None:
            for child in old:
                new.append(child)  # 굵게·크기 등 이 런에 따로 준 서식이 우선
            run._r.remove(old)
        run._r.insert(0, new)
    return run


def set_text(p, text):
    """문단의 첫 텍스트 런만 바꾸고 나머지 텍스트 런은 비운다(각주 참조 런은 남긴다)."""
    runs = [r for r in p.runs if r._r.find(qn("w:t")) is not None]
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def inline(p, text, size=None):
    """**굵게**, *기울임*, `코드`만 처리."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
        elif part.startswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)
        styled(r)
        if part.startswith("`"):
            r.font.name = "Courier New"
    return p


def borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "999999")
        b.append(e)
    tbl_pr.append(b)


def add_table(d, rows):
    cells = [[c.strip().replace("\x00", "|") for c in r.replace("\\|", "\x00").strip().strip("|").split("|")]
             for r in rows if not re.match(r"^\|[-| ]+\|$", r.strip())]
    n = max(len(r) for r in cells)
    t = d.add_table(rows=len(cells), cols=n)
    borders(t)
    size = 7 if n > 7 else 8
    # 열 폭을 내용 길이에 맞춰 본문 폭(16.5cm)에 나눈다. 기본값(균등·좁음)이면 글자가 줄마다 꺾여 쪽수가 늘어난다
    t.autofit = False
    longest = [max(len(re.sub(r"\*\*|`", "", r[j])) if j < len(r) else 0 for r in cells) for j in range(n)]
    weight = [min(max(x, 8), 60) ** 0.75 for x in longest]
    widths = [Cm(16.5 * w / sum(weight)) for w in weight]
    for i, r in enumerate(cells):
        for j in range(n):
            t.cell(i, j).width = widths[j]
    # 셀 폭만으로는 LibreOffice·Docs가 무시하므로 표 폭·고정 배치·격자 열 폭도 같이 쓴다
    pr = t._tbl.tblPr
    for tag, attrs in (("w:tblW", {"w:w": str(sum(w.twips for w in widths)), "w:type": "dxa"}), ("w:tblLayout", {"w:type": "fixed"})):
        old = pr.find(qn(tag))
        if old is not None:
            pr.remove(old)
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        pr.append(e)
    for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w.twips))
    for i, r in enumerate(cells):
        for j in range(n):
            p = t.cell(i, j).paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            inline(p, r[j] if j < len(r) else "", size)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def body(d, md):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1; continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(d, block); continue
        m = re.match(r"!\[[^\]]*\]\((figs/(fig\d)[^)]*)\)", line)
        if m:
            d.add_picture(str(HERE / m.group(1)), width=Cm(FIG_WIDTH[m.group(2)]))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if line.startswith("### "):
            styled(d.add_paragraph(style="Heading 3").add_run(line[4:]), "Heading 3")
        elif line.startswith("## "):
            styled(d.add_paragraph(style="Heading 2").add_run(line[3:]), "Heading 2")
        elif line.startswith("- "):
            p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.4)
            inline(p, "• " + line[2:])
        elif re.match(r"^\d+\. ", line):
            p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.5)
            inline(p, line, 9 if "arxiv.org" in line else None)
        elif line.startswith(("**Figure", "**Table")):
            inline(d.add_paragraph(), line, 9)
        else:
            inline(d.add_paragraph(), line)
        i += 1


def main():
    template, out = sys.argv[1], sys.argv[2]
    md = (HERE / "report.md").read_text()
    abstract = md.split("## Abstract", 1)[1].split("\n## ", 1)[0].strip()
    rest = "## " + md.split("## Abstract", 1)[1].split("\n## ", 1)[1]
    d = docx.Document(template)

    head = d.tables[0]
    set_text(next(p for p in head.rows[0].cells[0].paragraphs if "PROJECT TITLE" in p.text), TITLE)
    authors, abstract_box = head.rows[1].cells[0].tables
    grid = authors.rows[0].cells[0].tables[0]
    first = grid.rows[0].cells[0]
    for p in first.paragraphs:  # 이름과 소속은 줄바꿈이 든 런 하나다
        for r in p.runs:
            if "Author name" in r.text:
                r.text = f"{AUTHOR}\n{AFFIL}"  # python-docx가 \n을 줄바꿈(w:br)으로 쓴다
    # 저자 1명: 3x2 칸(각 1.2cm)을 칸 하나로 줄이고 표 폭 전체를 준다. 그대로 두면 이름이 글자 단위로 꺾인다
    tbl = grid._tbl
    for tr in tbl.findall(qn("w:tr"))[1:]:
        tbl.remove(tr)
    tr0 = tbl.find(qn("w:tr"))
    for tc in tr0.findall(qn("w:tc"))[1:]:
        tr0.remove(tc)
    cols = tbl.tblGrid.findall(qn("w:gridCol"))
    total = sum(int(float(c.get(qn("w:w")))) for c in cols)
    for c in cols[1:]:
        tbl.tblGrid.remove(c)
    cols[0].set(qn("w:w"), str(total))
    box = abstract_box.rows[0].cells[0]
    paras = [p for p in box.paragraphs if p.text.strip()]
    set_text(paras[0], abstract)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)

    # 안내 상자(두 번째 표)와 그 뒤 안내 문단을 전부 지우고, 마지막 sectPr만 남긴다
    for p in d.paragraphs:  # 지우기 전에 템플릿 런 서식을 떼어 둔다
        r = next((r for r in p.runs if r.text.strip()), None)
        kind = p.style.name if p.style.name in ("Heading 2", "Heading 3") else "body"
        if r is not None and kind not in RPR and r._r.find(qn("w:rPr")) is not None:
            RPR[kind] = copy.deepcopy(r._r.find(qn("w:rPr")))
    bd = d.element.body
    info = d.tables[1]._tbl
    started = False
    for el in list(bd.iterchildren()):
        if el is info:
            started = True
        if started and el.tag != qn("w:sectPr"):
            bd.remove(el)
    body(d, rest)
    d.save(out)
    print("저장:", out)


if __name__ == "__main__":
    main()
